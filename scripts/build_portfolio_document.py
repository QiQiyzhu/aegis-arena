"""Build the Chinese portfolio report from reviewed prose and native evidence.

Use the Python resolved by load_workspace_dependencies. The first DOCX creation
must follow the skill's one-time mark_artifact_operation_started.mjs command.
This script never calls that marker: repeated render iterations must not repeat it.

Draft preparation, without authoring an Office/PDF artifact:
    python build_portfolio_document.py --check-draft
    python build_portfolio_document.py --manifest-template D:/.../evidence.json

Final authoring, after real captures and validation are supplied:
    python build_portfolio_document.py --manifest D:/.../evidence.json \
        --artifact-operation-marked --qa-dir D:/.../document-qa \
        --renderer <documents-skill>/render_docx.py \
        --soffice <approved-isolated-or-bundled>/program/soffice.exe \
        --poppler-dir <bundled-poppler>/Library/bin

The canonical skill renderer produces PDF and page PNGs. Every PNG must still be
visually inspected before delivery; the machine checks cannot assert visual QA.
Native capture images are embedded unchanged, with aspect ratio preserved.
Only the two explicitly labelled design diagrams are drawn by this script.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
from datetime import datetime, timezone

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/portfolio-v1.5-design.md"
OUTPUT = ROOT / "outputs/portfolio-v1.5/documents"
FIELDS = ("build_summary", "native_summary", "video_summary", "validation_limits", "delivery_summary")
FIGURES = ("hero", "energy", "upgrade", "before", "after")
PAGES = ("overview", "loop", "economy", "stages", "upgrades", "ai", "visuals", "evidence", "references")
INLINE = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\(https://[^)]+\))")
PLACEHOLDER = re.compile(r"\{\{([a-z_]+)\}\}")


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def strict_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"Duplicate JSON key: {key}")
        result[key] = value
    return result


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8-sig"), object_pairs_hook=strict_object,
                      parse_constant=lambda value: (_ for _ in ()).throw(ValueError(value)))


def parse_source(text: str):
    parts = re.split(r"<!-- page: ([a-z]+) -->", text)
    if len(parts) != 19 or tuple(parts[1::2]) != PAGES:
        raise ValueError("Source must retain the nine reviewed page sections")
    prefix = parts[0].strip()
    pages = [(parts[i], parts[i + 1].strip()) for i in range(1, len(parts), 2)]
    for line in text.splitlines():
        if line.startswith("#"):
            heading = line.lstrip("# ")
            if not all(char.isalnum() or char.isspace() for char in heading):
                raise ValueError(f"Heading punctuation is not allowed: {heading}")
    names = set(PLACEHOLDER.findall(text))
    if names != set(FIELDS):
        raise ValueError(f"Unexpected evidence fields: {sorted(names)}")
    return prefix, pages


def template():
    return {
        "schemaVersion": 1,
        "reviewed": False,
        "fields": {name: "REPLACE_WITH_VERIFIED_CHINESE_TEXT" for name in FIELDS},
        "figures": {name: {
            "path": "D:/REPLACE/native-capture.png",
            "caption": "REPLACE_WITH_FACTUAL_CAPTION",
            "version": "REPLACE_WITH_BUILD_LABEL",
            "inputMode": "REPLACE_WITH_AUTOMATIC_OR_MANUAL_INPUT_SCOPE",
            "captureKind": "native",
            "sourceEvidence": "REPLACE_WITH_CAPTURE_OR_VIDEO_RECORD",
        } for name in FIGURES},
        "evidenceFiles": [],
        "notes": "Set reviewed true only after checking prose against source and actual evidence."
    }


def validate_manifest(manifest, source_text: str):
    from PIL import Image
    if manifest.get("schemaVersion") != 1 or manifest.get("reviewed") is not True:
        raise ValueError("The final evidence manifest has not been reviewed")
    fields = manifest.get("fields", {})
    if set(fields) != set(FIELDS):
        raise ValueError("Manifest fields do not match the reviewed draft")
    for name, value in fields.items():
        if not isinstance(value, str) or not 5 <= len(value) <= 900:
            raise ValueError(f"Evidence field length: {name}")
        if any(token in value for token in ("REPLACE", "待填", "TODO", "{{", "}}")) or "\n" in value:
            raise ValueError(f"Unresolved evidence field: {name}")
    figures = manifest.get("figures", {})
    if set(figures) != set(FIGURES):
        raise ValueError("Exactly the five reviewed native image slots are required")
    image_records = []
    for name, figure in figures.items():
        if figure.get("captureKind") != "native":
            raise ValueError(f"{name} must be a native capture, not a design diagram")
        for key in ("path", "caption", "version", "inputMode", "sourceEvidence"):
            value = figure.get(key)
            if not isinstance(value, str) or not value.strip() or "REPLACE" in value:
                raise ValueError(f"Missing figure provenance: {name}.{key}")
        path = Path(figure["path"])
        if not path.is_absolute() or not path.is_file():
            raise ValueError(f"Native image must exist at an absolute path: {name}")
        if path.suffix.lower() not in {".png", ".jpg", ".jpeg"}:
            raise ValueError(f"Unsupported native image type: {name}")
        with Image.open(path) as im:
            im.verify()
        with Image.open(path) as im:
            width, height = im.size
        if width < 960 or height < 540:
            raise ValueError(f"Image is too small for document review: {name}")
        image_records.append({"slot": name, "path": str(path), "width": width,
                              "height": height, "sha256": digest(path), **figure})
    evidence = []
    for item in manifest.get("evidenceFiles", []):
        path = Path(item)
        if not path.is_absolute() or not path.is_file():
            raise ValueError(f"Missing evidence file: {item}")
        evidence.append({"path": str(path), "sha256": digest(path)})
    if not evidence:
        raise ValueError("Bind at least one actual capture/validation report")
    rendered = PLACEHOLDER.sub(lambda match: fields[match.group(1)], source_text)
    remaining = re.sub(r"\{\{(?:figure|diagram):[a-z]+\}\}", "", rendered)
    if "{{" in remaining or "}}" in remaining:
        raise ValueError("Unresolved document placeholder")
    return rendered, image_records, evidence


def load_font(size: int):
    from PIL import ImageFont
    for candidate in (Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/msyh.ttc",
                      Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts/simhei.ttf"):
        if candidate.is_file():
            return ImageFont.truetype(str(candidate), size)
    raise FileNotFoundError("A verified Chinese font is required for the design diagrams")


def draw_diagrams(directory: Path):
    """Diagrams only: no edits to any screenshot pixels."""
    from PIL import Image, ImageDraw
    ink, accent, muted = "#142B39", "#177F83", "#667782"
    font, small = load_font(27), load_font(23)

    def arrow(draw, start, end, color=accent, width=4):
        draw.line([start, end], fill=color, width=width)
        angle = math.atan2(end[1] - start[1], end[0] - start[0])
        tip = [(end[0] - 16 * math.cos(angle + off), end[1] - 16 * math.sin(angle + off))
               for off in (-0.4, 0.4)]
        draw.polygon([end, *tip], fill=color)

    def label(draw, xy, text, face=font, color=ink):
        draw.multiline_text(xy, text, font=face, fill=color, anchor="mm", align="center", spacing=6)

    canvas = Image.new("RGB", (1440, 475), "white")
    draw = ImageDraw.Draw(canvas)
    nodes = [(90, 85, 400, 185, "准备\n选择压力档位"),
             (550, 85, 860, 185, "争夺目标与交火\n击杀获取能量"),
             (1010, 85, 1320, 185, "前两阶段完成\n奖励与恢复"),
             (1010, 290, 1320, 390, "选择升级\n暂停后进入下一阶段"),
             (550, 290, 860, 390, "撤离成功或失败\n结算本次行动"),
             (90, 290, 400, 390, "结算与重开\n资源和升级归零")]
    for x1, y1, x2, y2, text in nodes:
        draw.rounded_rectangle((x1, y1, x2, y2), radius=10, outline="#90ADB8", width=2)
        label(draw, ((x1+x2)/2, (y1+y2)/2), text)
    arrow(draw, (410,135),(535,135)); arrow(draw,(870,135),(995,135))
    arrow(draw,(1165,195),(1165,275))
    arrow(draw,(535,340),(415,340)); arrow(draw,(245,280),(245,200))
    draw.line([(1000,340),(942,340),(942,232),(805,232),(805,194)], fill=accent, width=4)
    arrow(draw,(805,232),(805,194))
    label(draw,(1063,236),"继续行动",small,muted)
    arrow(draw,(660,195),(660,275),color=ink)
    label(draw,(485,236),"最终撤离／死亡超时",small,muted)
    path_loop = directory / "design-loop.png"
    canvas.save(path_loop)

    canvas = Image.new("RGB", (1440, 710), "white")
    draw = ImageDraw.Draw(canvas)
    # World XY coordinates from SetObjectiveLocation and the arena cover geometry.
    def project(x, y):
        return 180 + (x + 2000) * 0.245, 72 + (1500 - y) * 0.18
    left, top = project(-2000, 1500)
    right, bottom = project(2000, -1500)
    draw.rectangle((left,top,right,bottom), outline="#96ADB7", width=3)
    for x, y, size in ((-300,400,360),(300,-400,360),(600,600,280),(-700,-600,280)):
        x1,y1=project(x-size/2,y+size/2); x2,y2=project(x+size/2,y-size/2)
        draw.rectangle((x1,y1,x2,y2), fill="#D8E0E4", outline="#8798A1", width=2)
    points = [(-500,-1000,"1 中继"),(850,-850,"2A 传输"),(0,1000,"2B 传输"),(-1250,0,"3 撤离")]
    for x,y,name in points:
        px,py=project(x,y)
        draw.ellipse((px-260*.245,py-260*.18,px+260*.245,py+260*.18),
                     fill="#E8F5F3",outline=accent,width=4)
        label(draw,(px,py),name,small)
    # Only the compulsory relay order is connected, never a suggested path.
    p1=project(850,-850); p2=project(0,1000)
    dx,dy=p2[0]-p1[0],p2[1]-p1[1]
    for i in range(2,17,2):
        a=i/20;b=(i+1)/20
        draw.line([(p1[0]+dx*a,p1[1]+dy*a),(p1[0]+dx*b,p1[1]+dy*b)],fill=accent,width=3)
    arrow(draw,(p1[0]+dx*.8,p1[1]+dy*.8),(p1[0]+dx*.87,p1[1]+dy*.87),width=3)
    label(draw,(1245,130),"灰色方块\n原有掩体",small,muted)
    label(draw,(1245,250),"绿色圈\n目标范围",small,muted)
    label(draw,(1245,390),"虚线\n2A 完成后\n激活 2B",small,muted)
    label(draw,(710,663),"设计示意  世界 XY 俯视  非实机截图  不代表导航路线",small,muted)
    path_arena=directory/"design-arena.png"
    canvas.save(path_arena)
    return {"loop": (path_loop,"核心循环设计示意。框线表示状态，箭头表示推进与重开关系，非实机截图。"),
            "arena": (path_arena,"地形与目标设计示意。目标坐标来自运行规则，灰块为原有掩体；虚线仅表示先后关系，非导航路径或实机截图。")}


def set_font(run, size=None, bold=None, color="000000"):
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    run.font.name="Microsoft YaHei"
    fonts=run._element.get_or_add_rPr().rFonts
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:"+key),None)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:"+key),"Microsoft YaHei")
    if size is not None: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    run.font.color.rgb=RGBColor.from_string(color)


def inline(paragraph, text, size=None, color="000000"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith("**") and part.endswith("**"):
            set_font(paragraph.add_run(part[2:-2]),size,True,color)
        elif part.startswith("["):
            match=re.fullmatch(r"\[([^\]]+)\]\((https://[^)]+)\)",part)
            if not match: set_font(paragraph.add_run(part),size,color=color); continue
            title,url=match.groups()
            link=OxmlElement("w:hyperlink")
            link.set(qn("r:id"),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
            run=paragraph.add_run(title);set_font(run,size,color="17626A")
            link.append(run._r);paragraph._p.append(link)
        else:
            set_font(paragraph.add_run(part),size,color=color)


def style_document(document):
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    section=document.sections[0]
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=Inches(.66);section.bottom_margin=Inches(.62)
    section.left_margin=section.right_margin=Inches(.72)
    section.footer_distance=Inches(.30)
    for name,size in (("Normal",11),("Title",25),("Subtitle",11),("Heading 1",18),
                      ("Heading 2",12.5),("Caption",9.2)):
        style=document.styles[name]
        style.font.name="Microsoft YaHei";style.font.size=Pt(size)
        style.font.color.rgb=RGBColor(0,0,0)
        style.font.bold=name.startswith("Heading")
        style.font.italic=False
        fonts=style._element.get_or_add_rPr().rFonts
        for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            fonts.attrib.pop(qn("w:"+key),None)
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:"+key),"Microsoft YaHei")
        props=style.paragraph_format
        props.space_after=Pt(7);props.line_spacing=1.15
        if name.startswith("Heading"):
            props.space_before=Pt(10);props.keep_with_next=True
        if name=="Caption":
            props.space_before=Pt(3);props.space_after=Pt(9);props.line_spacing=1.08
        # Prevent a template's inherited decorative rule from surviving.
        ppr=style._element.find(qn("w:pPr"))
        if ppr is not None:
            for element in list(ppr):
                if element.tag in {qn("w:pBdr"),qn("w:shd")}:
                    ppr.remove(element)
    footer=section.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    set_font(footer.add_run("Aegis Arena  |  "),8,color="5B6770")
    field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE")
    footer._p.append(field)
    document.core_properties.title="Aegis Arena 系统设计作品说明"
    document.core_properties.subject="系统与综合策划作品集 1.5"
    document.core_properties.author="Aegis Arena 项目协作"
    document.core_properties.keywords="系统策划,资源循环,目标,成长,验证"
    document.core_properties.comments=""


def table(document, rows, page):
    from docx.shared import Inches, Pt
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    count=len(rows[0]); widths=([1.30,5.76] if count==2 else [1.32,1.75,3.99])
    if page=="evidence":widths=[2.30,2.40,2.36]
    if page=="loop":widths=[1.2,2.7,3.16]
    result=document.add_table(rows=0,cols=count);result.autofit=False
    result.alignment=WD_TABLE_ALIGNMENT.CENTER
    for column,width in zip(result.columns,widths):column.width=Inches(width)
    for index,items in enumerate(rows):
        if len(items)!=count:raise ValueError("Malformed markdown table")
        row=result.add_row()
        cant=OxmlElement("w:cantSplit");row._tr.get_or_add_trPr().append(cant)
        if index==0:
            header=OxmlElement("w:tblHeader");row._tr.get_or_add_trPr().append(header)
        for ci,(cell,value) in enumerate(zip(row.cells,items)):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            props=cell._tc.get_or_add_tcPr()
            borders=OxmlElement("w:tcBorders")
            for side in ("top","left","bottom","right"):
                border=OxmlElement(f"w:{side}")
                for key,val in (("val","single"),("sz","4"),("color","D9D9D9")):
                    border.set(qn(f"w:{key}"),val)
                borders.append(border)
            props.append(borders)
            margins=OxmlElement("w:tcMar")
            for side,value_margin in (("top",85),("bottom",85),("left",115),("right",115)):
                margin=OxmlElement(f"w:{side}");margin.set(qn("w:w"),str(value_margin));margin.set(qn("w:type"),"dxa");margins.append(margin)
            props.append(margins)
            shading=OxmlElement("w:shd");shading.set(qn("w:fill"),"193E50" if index==0 else ("F0F4F6" if index%2==0 else "FFFFFF"));props.append(shading)
            paragraph=cell.paragraphs[0];paragraph.paragraph_format.space_after=Pt(0)
            paragraph.paragraph_format.line_spacing=1.13
            paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER if ci==0 else WD_ALIGN_PARAGRAPH.LEFT
            inline(paragraph,value,size=9.8,color="FFFFFF" if index==0 else "152A35")
            if index==0:
                for run in paragraph.runs:run.bold=True
    spacer=document.add_paragraph();spacer.paragraph_format.space_after=Pt(2);spacer.paragraph_format.space_before=Pt(0)
    spacer.paragraph_format.line_spacing=1;spacer.paragraph_format.line_spacing=Pt(2)


def picture(document,path: Path,caption: str,max_height: float,width: float=7.06):
    from PIL import Image
    from docx.shared import Inches,Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    with Image.open(path) as image:
        ratio=image.width/image.height
    width=min(width,max_height*ratio)
    paragraph=document.add_paragraph();paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next=True;paragraph.paragraph_format.space_after=Pt(0)
    shape=paragraph.add_run().add_picture(str(path),width=Inches(width))
    shape._inline.docPr.set("descr",caption)
    paragraph=document.add_paragraph(style="Caption")
    inline(paragraph,caption,size=9.2,color="465966")


def blocks(document,text,page,figures,diagrams,counter):
    lines=text.splitlines();index=0
    while index<len(lines):
        line=lines[index].strip()
        if not line:index+=1;continue
        if line.startswith("|"):
            rows=[]
            while index<len(lines) and lines[index].strip().startswith("|"):
                cells=[part.strip() for part in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?",cell) for cell in cells):rows.append(cells)
                index+=1
            table(document,rows,page);continue
        match=re.fullmatch(r"\{\{(figure|diagram):([a-z]+)\}\}",line)
        if match:
            kind,name=match.groups();counter[0]+=1
            if kind=="figure":
                item=figures[name]
                caption=f"图 {counter[0]}  {item['caption']}  {item['version']}；{item['inputMode']}。"
                height={"hero":2.8,"energy":2.4,"upgrade":2.45,"before":2.85,"after":2.85}[name]
                picture(document,Path(item["path"]),caption,height)
            else:
                path,caption=diagrams[name]
                picture(document,path,f"图 {counter[0]}  {caption}",2.02 if name=="loop" else 2.95)
            index+=1;continue
        if line.startswith("### "):
            document.add_heading(line[4:],level=2)
        elif line.startswith("## "):
            document.add_heading(line[3:],level=1)
        elif line.startswith("# "):
            document.add_paragraph(line[2:],style="Title")
        else:
            paragraph=document.add_paragraph()
            inline(paragraph,line)
        index+=1


def render(docx: Path, qa: Path, renderer: Path, soffice: Path, poppler: Path):
    if not renderer.is_file() or renderer.name!="render_docx.py":
        raise ValueError("Use the canonical documents skill render_docx.py")
    if not soffice.is_file() or soffice.name.lower()!="soffice.exe":
        raise ValueError("Supply the explicit bundled or task-approved isolated soffice.exe")
    if not (poppler/"pdftoppm.exe").is_file() or not (poppler/"pdfinfo.exe").is_file():
        raise ValueError("Supply the bundled Poppler binary directory")
    env=os.environ.copy()
    # The renderer calls shutil.which. This child PATH deterministically selects
    # only the explicit approved copy; no global PATH or file association changes.
    env["PATH"]=str(soffice.parent)+os.pathsep+str(poppler)+os.pathsep+env.get("PATH","")
    if Path(shutil.which("soffice.exe",path=env["PATH"])).resolve()!=soffice.resolve():
        raise ValueError("The renderer would select a different LibreOffice")
    temp=qa/"temp";temp.mkdir();env["TEMP"]=str(temp);env["TMP"]=str(temp)
    output=qa/"render"
    command=[sys.executable,str(renderer),str(docx),"--output_dir",str(output),"--emit_pdf","--dpi","150","--verbose"]
    with (qa/"render-command.json").open("w",encoding="utf-8") as stream:
        json.dump({"command":command,"soffice":str(soffice),"poppler":str(poppler)},stream,ensure_ascii=False,indent=2)
    run=subprocess.run(command,cwd=str(ROOT),env=env,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True,
                       encoding="utf-8",errors="replace",timeout=180)
    (qa/"renderer.log").write_text(run.stdout,encoding="utf-8")
    if run.returncode:raise RuntimeError(f"Document rendering failed. See {qa/'renderer.log'}")
    pdf=output/(docx.stem+".pdf")
    if not pdf.is_file():raise RuntimeError("Renderer did not emit the requested PDF")
    from pypdf import PdfReader
    reader=PdfReader(pdf)
    if not 8<=len(reader.pages)<=10:raise ValueError(f"Unexpected rendered page count: {len(reader.pages)}")
    pngs=list(output.glob("page-*.png"))
    if len(pngs)!=len(reader.pages):raise ValueError("Not every PDF page has a QA image")
    for index,page in enumerate(reader.pages,1):
        if len((page.extract_text() or "").strip())<30:raise ValueError(f"Unexpected blank page: {index}")
    return pdf,len(reader.pages),[str(path) for path in pngs]


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__,formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--source",type=Path,default=SOURCE)
    parser.add_argument("--output",type=Path,default=OUTPUT)
    parser.add_argument("--manifest",type=Path)
    parser.add_argument("--manifest-template",type=Path)
    parser.add_argument("--check-draft",action="store_true")
    parser.add_argument("--check-only",action="store_true")
    parser.add_argument("--artifact-operation-marked",action="store_true")
    parser.add_argument("--qa-dir",type=Path)
    parser.add_argument("--renderer",type=Path)
    parser.add_argument("--soffice",type=Path)
    parser.add_argument("--poppler-dir",type=Path)
    args=parser.parse_args(argv)
    text=args.source.read_text(encoding="utf-8");parse_source(text)
    if args.check_draft:
        print(json.dumps({"pagesPlanned":9,"figuresRequired":FIGURES,"diagrams":["loop","arena"],"evidenceFields":FIELDS,"artifactCreated":False}));return
    if args.manifest_template:
        if args.manifest_template.exists():raise FileExistsError(args.manifest_template)
        args.manifest_template.parent.mkdir(parents=True,exist_ok=True)
        args.manifest_template.write_text(json.dumps(template(),ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
        print(args.manifest_template);return
    if not args.manifest:parser.error("--manifest is required for final evidence")
    manifest=read_json(args.manifest)
    rendered,image_records,evidence=validate_manifest(manifest,text)
    if args.check_only:print("Evidence structure and image inputs validated; no artifact created.");return
    if not args.artifact_operation_marked:parser.error("Run the skill's one-time marker before the first DOCX create, then supply --artifact-operation-marked")
    if not all((args.qa_dir,args.renderer,args.soffice,args.poppler_dir)):
        parser.error("Final authoring requires --qa-dir, --renderer, --soffice, --poppler-dir")
    if args.qa_dir.exists():raise FileExistsError("Choose a fresh QA directory for each render iteration")
    args.qa_dir.mkdir(parents=True);args.output.mkdir(parents=True,exist_ok=True)
    from docx import Document
    document=Document();style_document(document)
    diagrams=draw_diagrams(args.qa_dir)
    # Parse after placeholder substitution; headings and page markers are unchanged.
    parts=re.split(r"<!-- page: ([a-z]+) -->",rendered)
    counter=[0];blocks(document,parts[0],"overview",manifest["figures"],diagrams,counter)
    for index in range(1,len(parts),2):
        # A dedicated break paragraph can itself overflow a full preceding page,
        # creating a blank page in LibreOffice. Break on the next heading instead.
        page_start=len(document.paragraphs)
        blocks(document,parts[index+1],parts[index],manifest["figures"],diagrams,counter)
        if index>1:document.paragraphs[page_start].paragraph_format.page_break_before=True
    docx=args.output/"Aegis_Arena_System_Design_v1.5.docx"
    document.save(docx)
    pdf,pages,pngs=render(docx,args.qa_dir,args.renderer,args.soffice,args.poppler_dir)
    final_pdf=args.output/(docx.stem+".pdf");shutil.copy2(pdf,final_pdf)
    for item in image_records:
        if digest(Path(item["path"]))!=item["sha256"]:raise RuntimeError("Native image changed while authoring")
    record={"createdUtc":datetime.now(timezone.utc).isoformat(),"python":sys.executable,
            "source":{"path":str(args.source),"sha256":digest(args.source)},
            "manifest":{"path":str(args.manifest),"sha256":digest(args.manifest)},
            "nativeImages":image_records,"evidenceFiles":evidence,"pageCount":pages,"pageImages":pngs,
            "designDiagrams":[{"path":str(path),"sha256":digest(path),"kind":"design_diagram"} for path,_ in diagrams.values()],
            "outputs":[{"path":str(path),"sha256":digest(path),"bytes":path.stat().st_size} for path in (docx,final_pdf)],
            "visualReviewRequired":True,"visualReviewCompleted":False}
    (args.qa_dir/"build-record.json").write_text(json.dumps(record,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"docx":str(docx),"pdf":str(final_pdf),"pages":pages,"qa":str(args.qa_dir),"requiresEveryPageVisualReview":True},ensure_ascii=False))


if __name__=="__main__":
    main()
