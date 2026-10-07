"""Prepare or render the v2.1 Chinese design report from reviewed native evidence.

No artifact is created by --check-draft, --manifest-template or --check-only.
Formal authoring requires the documents skill's one-time marker, an external
reviewed manifest, a fresh QA directory and its canonical render_docx.py.
Every rendered page still requires manual visual inspection before delivery.
The v1.5 helper is imported for typography/rendering only and never modified.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys

import build_portfolio_document as layout

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/portfolio-v2.1-design.md"
OUTPUT = ROOT / "outputs/portfolio-v2.1/documents"
FIELDS = ("build_summary", "run_summary", "native_summary", "audio_validation", "vfx_validation", "video_summary", "validation_limits", "delivery_summary")
FIGURES = ("hero", "energy", "upgrade", "shot", "charge", "diagnostic", "after")
PAGES = ("overview", "loop", "economy", "objectives", "combat", "music", "ai", "history", "evidence", "video", "references")
DIAGRAMS = ("loop", "route", "audio")
HEIGHTS = {"hero": 2.3, "energy": 1.85, "upgrade": 1.6, "shot": 1.7,
           "charge": 1.7, "diagnostic": 2.65, "after": 2.6}
TEXT_FIELD = re.compile(r"\{\{([a-z_]+)\}\}")
MEDIA_FIELD = re.compile(r"\{\{(figure|diagram):([a-z]+)\}\}")
DRAFT_BLOCK = re.compile(r"<!-- draft-only -->.*?<!-- /draft-only -->", re.S)


def require(condition, message):
    if not condition:
        raise ValueError(message)


def parse_source(text):
    text = DRAFT_BLOCK.sub("", text)
    parts = re.split(r"<!-- page: ([a-z]+) -->", text)
    require(tuple(parts[1::2]) == PAGES, "The reviewed eleven-page outline must be retained")
    for line in text.splitlines():
        if line.startswith("#"):
            require(all(c.isalnum() or c.isspace() for c in line.lstrip("# ")),
                    f"Heading punctuation is not allowed: {line}")
    require(set(TEXT_FIELD.findall(text)) == set(FIELDS), "Evidence fields do not match the v2.1 draft")
    media = MEDIA_FIELD.findall(text)
    require(sorted(n for kind, n in media if kind == "figure") == sorted(FIGURES),
            "Require the seven native figure slots exactly once")
    require(sorted(n for kind, n in media if kind == "diagram") == sorted(DIAGRAMS),
            "Require the three labelled design diagrams exactly once")
    return text


def template():
    return {"schemaVersion": 3, "reviewed": False, "sourceSha256": "REPLACE_WITH_FINAL_SOURCE_SHA256",
            "fields": {name: "待验证：填写与实际证据一致的中文范围与结果" for name in FIELDS},
            "figures": {name: {"path": "D:/REPLACE/native-frame.png",
                               "caption": "REPLACE_WITH_FACTUAL_CAPTION",
                               "version": "v2.1 / REPLACE_WITH_SOURCE_PREFIX",
                               "inputMode": "REPLACE_WITH_ACTUAL_INPUT_SCOPE",
                               "captureKind": "native",
                               "sourceEvidence": "D:/REPLACE/capture-provenance.json"} for name in FIGURES},
            "evidenceFiles": [],
            "notes": "Only set reviewed true after checking prose, frozen source, native gates and complete video evidence."}


def absolute_file(value, label):
    require(isinstance(value, str) and value.strip(), f"Missing {label}")
    path = Path(value)
    require(path.is_absolute() and path.is_file(), f"{label} must be an existing absolute path")
    return path.resolve()


def validate_manifest(manifest, source_text):
    from PIL import Image
    require(isinstance(manifest, dict), "Manifest must be an object")
    require(type(manifest.get("schemaVersion")) is int and manifest["schemaVersion"] == 3,
            "Expected v2.1 evidence schema")
    require(manifest.get("reviewed") is True, "An explicit human/agent evidence review is required")
    source_sha = manifest.get("sourceSha256", "")
    require(isinstance(source_sha, str) and re.fullmatch(r"[0-9a-f]{64}", source_sha),
            "Bind the final frozen source SHA256")
    fields = manifest.get("fields")
    require(isinstance(fields, dict) and set(fields) == set(FIELDS), "Evidence field names mismatch")
    for name, value in fields.items():
        require(isinstance(value, str) and 5 <= len(value) <= 1000, f"Invalid evidence prose: {name}")
        require(not any(token in value for token in ("REPLACE", "TODO", "待填", "{{", "}}", "\n", "\r")),
                f"Unresolved evidence prose: {name}")
    require(source_sha[:12] in fields["build_summary"], "Build summary must identify the frozen source")

    raw_evidence = manifest.get("evidenceFiles")
    require(isinstance(raw_evidence, list) and bool(raw_evidence), "Bind real reports, not a prose-only manifest")
    evidence = []
    evidence_paths = set()
    for value in raw_evidence:
        path = absolute_file(value, "evidence file")
        require(path not in evidence_paths, "Duplicate evidence file")
        evidence_paths.add(path)
        evidence.append({"path": str(path), "sha256": layout.digest(path), "bytes": path.stat().st_size})

    figures = manifest.get("figures")
    require(isinstance(figures, dict) and set(figures) == set(FIGURES), "Native figure slots mismatch")
    records = []
    for name, figure in figures.items():
        require(isinstance(figure, dict) and figure.get("captureKind") == "native", f"{name} must be native")
        for key in ("caption", "version", "inputMode"):
            value = figure.get(key)
            require(isinstance(value, str) and value.strip() and "REPLACE" not in value,
                    f"Missing figure description: {name}.{key}")
        require("2.1" in figure["version"] and source_sha[:12] in figure["version"],
                f"{name} must identify the current v2.1 build")
        path = absolute_file(figure.get("path"), f"{name} screenshot")
        proof = absolute_file(figure.get("sourceEvidence"), f"{name} source report")
        require(proof in evidence_paths, f"{name} source report must be bound in evidenceFiles")
        require(path.suffix.lower() in (".png", ".jpg", ".jpeg"), "Unsupported native image type")
        with Image.open(path) as im:
            im.verify()
        with Image.open(path) as im:
            width, height = im.size
        require(width >= 960 and height >= 540, f"{name} image is too small")
        records.append({**figure, "slot": name, "path": str(path), "sourceEvidence": str(proof),
                        "width": width, "height": height, "sha256": layout.digest(path)})
    rendered = TEXT_FIELD.sub(lambda m: fields[m.group(1)], parse_source(source_text))
    residue = MEDIA_FIELD.sub("", rendered)
    require("{{" not in residue and "}}" not in residue, "Unresolved template token")
    return rendered, records, evidence


def draw_diagrams(directory):
    """Original rule schematics only; never edit screenshot pixels."""
    from PIL import Image, ImageDraw
    font = layout.load_font(27)
    small = layout.load_font(22)
    ink, accent, muted = "#142B39", "#177F83", "#667782"

    def node(draw, box, text, face=font):
        draw.rounded_rectangle(box, radius=10, outline="#90ADB8", width=2)
        draw.multiline_text(((box[0]+box[2])/2, (box[1]+box[3])/2), text, font=face,
                            fill=ink, anchor="mm", align="center", spacing=6)

    def arrow(draw, x1, y, x2):
        draw.line((x1, y, x2, y), fill=accent, width=4)
        draw.polygon(((x2,y),(x2-14,y-7),(x2-14,y+7)),fill=accent)

    im = Image.new("RGB", (1440, 440), "white")
    d = ImageDraw.Draw(im)
    node(d, (45,65,335,160), "准备与选择顺序")
    node(d, (440,65,840,160), "交火与目标推进\n击杀和阶段奖励供能")
    node(d, (945,65,1380,160), "升级后继续／最终撤离\n死亡或超时则失败")
    arrow(d, 345,112,425);arrow(d,850,112,930)
    boxes=[(45,260,345,355,"RMB 蓄能\n时间与能量换排线"),
           (395,260,695,355,"Q 脉冲\n能量换即时解围"),
           (745,260,1045,355,"E 修复\n按实际恢复付费"),
           (1095,260,1395,355,"F 超频\n能量换目标时间")]
    for x1,y1,x2,y2,text in boxes:
        node(d,(x1,y1,x2,y2),text,small)
    d.line((640,170,640,215),fill=accent,width=4)
    d.line((195,215,1245,215),fill=accent,width=3)
    for x in (195,545,895,1245):d.line((x,215,x,248),fill=accent,width=3)
    d.text((720,405),"同一能量账本约束四项用途  重开回到初始状态",font=small,fill=muted,anchor="mm")
    loop=directory/"v21-design-loop.png";im.save(loop)

    im=Image.new("RGB",(1440,330),"white");d=ImageDraw.Draw(im)
    for row, labels in enumerate((("首段 8 秒","北点 10 秒","东点 10 秒","撤离 4 秒"),
                                   ("首段 8 秒","东点 10 秒","北点 10 秒","撤离 4 秒"))):
        y=35+row*125
        for index,label in enumerate(labels):
            x=35+index*360
            node(d,(x,y,x+275,y+85),label)
            if index<3:arrow(d,x+285,y+42,x+340)
    d.text((720,305),"两种目标顺序设计示意  非地图截图  箭头不代表导航路线",font=small,fill=muted,anchor="mm")
    route=directory/"v21-design-order.png";im.save(route)
    im=Image.new("RGB",(1440,370),"white");d=ImageDraw.Draw(im)
    for index,label in enumerate(("公开阶段与菜单", "选择对应音乐", "最多两轨交叉淡化", "音频组件实际播放")):
        x=25+index*358
        node(d,(x,40,x+285,140),label,small)
        if index<3:arrow(d,x+292,90,x+340)
    node(d,(180,215,690,305),"M 静音／恢复  暂停也可操作",small)
    node(d,(765,215,1280,305),"重开与退出  停止旧组件",small)
    d.text((720,346),"音频状态设计示意  不读取敌人隐藏状态  非实机截图",font=small,fill=muted,anchor="mm")
    audio=directory/"v21-design-audio.png";im.save(audio)
    return {"loop":(loop,"核心循环与能量用途设计示意，非实机截图。"),
            "route":(route,"目标顺序设计示意，时间为单人基础条件；不表示实际通关耗时。"),
            "audio":(audio,"音乐状态与生命周期设计示意，非实机截图。")}


def blocks(document, text, page, figures, diagrams, counter):
    lines=text.splitlines();index=0
    while index<len(lines):
        line=lines[index].strip()
        if not line:index+=1;continue
        if line.startswith("|"):
            rows=[]
            while index<len(lines) and lines[index].strip().startswith("|"):
                cells=[v.strip() for v in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?",c) for c in cells):rows.append(cells)
                index+=1
            layout.table(document,rows,page);continue
        media=MEDIA_FIELD.fullmatch(line)
        if media:
            kind,name=media.groups();counter[0]+=1
            if kind=="figure":
                item=figures[name]
                caption=f"图 {counter[0]}  {item['caption']}  {item['version']}；{item['inputMode']}。"
                layout.picture(document,Path(item["path"]),caption,HEIGHTS[name])
            else:
                path,caption=diagrams[name]
                layout.picture(document,path,f"图 {counter[0]}  {caption}",{"loop":1.75,"route":1.25,"audio":1.4}[name])
        elif line.startswith("### "):document.add_heading(line[4:],level=2)
        elif line.startswith("## "):document.add_heading(line[3:],level=1)
        elif line.startswith("# "):document.add_paragraph(line[2:],style="Title")
        else:layout.inline(document.add_paragraph(),line)
        index+=1


def render_document(docx, qa, renderer, soffice, poppler):
    """Use the canonical renderer with the v2.1 eleven-page contract."""
    require(renderer.is_file() and renderer.name=="render_docx.py", "Canonical renderer required")
    require(soffice.is_file() and soffice.name.lower()=="soffice.exe", "Explicit LibreOffice required")
    require(all((poppler/name).is_file() for name in ("pdftoppm.exe","pdfinfo.exe")), "Bundled Poppler required")
    env=os.environ.copy()
    env["PATH"]=str(soffice.parent)+os.pathsep+str(poppler)+os.pathsep+env.get("PATH","")
    require(Path(shutil.which("soffice.exe",path=env["PATH"])).resolve()==soffice.resolve(), "Wrong renderer runtime")
    temp=qa/"temp";temp.mkdir();env["TEMP"]=str(temp);env["TMP"]=str(temp)
    output=qa/"render"
    command=[sys.executable,str(renderer),str(docx),"--output_dir",str(output),"--emit_pdf","--dpi","150","--verbose"]
    (qa/"render-command.json").write_text(json.dumps({"command":command,"soffice":str(soffice),"poppler":str(poppler)},indent=2),encoding="utf-8")
    result=subprocess.run(command,cwd=ROOT,env=env,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,
                          text=True,encoding="utf-8",errors="replace",timeout=180)
    (qa/"renderer.log").write_text(result.stdout,encoding="utf-8")
    require(result.returncode==0, "Document rendering failed; inspect renderer.log")
    pdf=output/(docx.stem+".pdf");require(pdf.is_file(), "Renderer emitted no PDF")
    from pypdf import PdfReader
    reader=PdfReader(pdf)
    pngs=sorted(output.glob("page-*.png"),key=lambda p:int(p.stem.split('-')[-1]))
    require(len(pngs)==len(reader.pages), "Missing page renders")
    require(all(len((p.extract_text() or "").strip())>=30 for p in reader.pages), "Unexpected blank page")
    return pdf,len(reader.pages),[str(p) for p in pngs]


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source",type=Path,default=SOURCE)
    parser.add_argument("--output",type=Path,default=OUTPUT)
    parser.add_argument("--manifest",type=Path)
    parser.add_argument("--manifest-template",type=Path)
    parser.add_argument("--check-draft",action="store_true")
    parser.add_argument("--check-only",action="store_true")
    parser.add_argument("--artifact-operation-marked",action="store_true")
    parser.add_argument("--qa-dir",type=Path)
    parser.add_argument("--renderer",type=Path)
    parser.add_argument("--soffice",type=Path)
    parser.add_argument("--poppler-dir",type=Path)
    args=parser.parse_args(argv)
    raw=args.source.read_text(encoding="utf-8");parse_source(raw)
    if args.check_draft:
        print(json.dumps({"pagesPlanned":11,"figuresRequired":FIGURES,"diagrams":DIAGRAMS,
                          "evidenceFields":FIELDS,"artifactCreated":False}));return
    if args.manifest_template:
        require(not args.manifest_template.exists(), "Manifest template must use a new path")
        args.manifest_template.parent.mkdir(parents=True,exist_ok=True)
        args.manifest_template.write_text(json.dumps(template(),ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
        print(args.manifest_template);return
    if not args.manifest:parser.error("--manifest is required for reviewed evidence")
    manifest=layout.read_json(args.manifest)
    rendered,images,evidence=validate_manifest(manifest,raw)
    if args.check_only:
        print("Reviewed evidence structure and image inputs validated; no artifact created.");return
    if not args.artifact_operation_marked:
        parser.error("Run the skill marker once before the first v2.1 artifact creation")
    if not all((args.qa_dir,args.renderer,args.soffice,args.poppler_dir)):
        parser.error("Formal authoring requires a fresh QA directory and canonical renderer dependencies")
    old_output=(ROOT/"outputs/portfolio-v1.5").resolve()
    require(args.output.resolve()!=old_output and old_output not in args.output.resolve().parents,
            "Never write v2.1 artifacts into the preserved v1.5 output")
    previous=(ROOT/"outputs/portfolio-v2").resolve()
    require(args.output.resolve()!=previous and previous not in args.output.resolve().parents,
            "Never overwrite preserved v2.0 artifacts")
    require(not args.qa_dir.exists(), "Use a fresh QA directory")
    watched=[args.source.resolve(),args.manifest.resolve(),Path(__file__).resolve(),Path(layout.__file__).resolve()]
    watched += [Path(r["path"]) for r in images+evidence]
    input_hashes={str(p):layout.digest(p) for p in watched}
    args.qa_dir.mkdir(parents=True);args.output.mkdir(parents=True,exist_ok=True)
    from docx import Document
    document=Document();layout.style_document(document)
    document.core_properties.subject="系统与综合策划作品集 2.1"
    diagrams=draw_diagrams(args.qa_dir)
    parts=re.split(r"<!-- page: ([a-z]+) -->",rendered);counter=[0]
    blocks(document,parts[0],"overview",manifest["figures"],diagrams,counter)
    for i in range(1,len(parts),2):
        first=len(document.paragraphs)
        blocks(document,parts[i+1],parts[i],manifest["figures"],diagrams,counter)
        if i>1:document.paragraphs[first].paragraph_format.page_break_before=True
    docx=args.output/"Aegis_Arena_System_Design_v2.1.docx";document.save(docx)
    pdf,pages,pngs=render_document(docx,args.qa_dir,args.renderer,args.soffice,args.poppler_dir)
    require(pages==len(PAGES), "Render must retain the eleven reviewed pages; adjust layout, not body font readability")
    final_pdf=args.output/(docx.stem+".pdf");shutil.copy2(pdf,final_pdf)
    for value,expected in input_hashes.items():
        require(layout.digest(Path(value))==expected, f"Input changed while rendering: {value}")
    record={"createdUtc":datetime.now(timezone.utc).isoformat(),"python":sys.executable,
            "schemaVersion":2,"sourceSha256":manifest["sourceSha256"],
            "source":{"path":str(args.source.resolve()),"sha256":layout.digest(args.source)},
            "manifest":{"path":str(args.manifest.resolve()),"sha256":layout.digest(args.manifest)},
            "inputHashes":input_hashes,"nativeImages":images,"evidenceFiles":evidence,
            "pageCount":pages,"pageImages":pngs,"visualReviewRequired":True,"visualReviewCompleted":False,
            "designDiagrams":[{"path":str(p),"sha256":layout.digest(p),"kind":"design_diagram"} for p,_ in diagrams.values()],
            "outputs":[{"path":str(p.resolve()),"sha256":layout.digest(p),"bytes":p.stat().st_size} for p in (docx,final_pdf)]}
    (args.qa_dir/"build-record.json").write_text(json.dumps(record,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"pages":pages,"docx":str(docx),"pdf":str(final_pdf),"qa":str(args.qa_dir),
                      "requiresEveryPageVisualReview":True},ensure_ascii=False))


if __name__=="__main__":
    main()
